"""Build lect-12-shaders-lighting.html with consolidated (theme-grouped) slides.

Text, paragraph levels, tables, hyperlinks and speaker notes are read from the
PPTX directly (the extractor JSON flattens levels and drops tables); images use
the extractor output in lect-12-shaders-lighting-content.json. The PLAN
maps every source slide (1-68) onto a smaller set of themed slides; see
docs/lecture-conversion-playbook.md.
"""
import html
import json
import re
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PAGES = ROOT / "site" / "pages"
PPTX_PATH = ROOT / "site" / "source-lecture-pptx" / "3DIMD Lecture 12 Unity Shaders Lighting  01.pptx"
JSON_PATH = PAGES / "lect-12-shaders-lighting-content.json"
OUT_PATH = PAGES / "lect-12-shaders-lighting.html"
FOOTER_SOURCE = PAGES / "lect-04a-report.html"

SLIDE_TITLE = "3D Interactive Media Development"
PAGE_TITLE = "3DIMD | Lecture 12 - Unity Lighting and Shaders"

images = {s["slide_number"]: s["images"] for s in json.load(JSON_PATH.open(encoding="utf-8"))}
prs = Presentation(str(PPTX_PATH))
SLIDES = {i: s for i, s in enumerate(prs.slides, 1)}


def _collect_lines(slide):
    out = []

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6:
                walk(shape.shapes)
                continue
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for seg in para.text.replace("\x0b", "\n").replace("\r", "\n").split("\n"):
                    lead = re.match(r"^[\t\u2003\xa0 ]*", seg).group(0)
                    text = re.sub(r"\s+", " ", seg.replace("\xa0", " ")).strip()
                    if text and text != SLIDE_TITLE:
                        out.append((para.level + lead.count("\t"), text))

    walk(slide.shapes)
    return out


LINES = {n: _collect_lines(s) for n, s in SLIDES.items()}


def lines(num):
    return LINES[num]


def esc(text):
    return html.escape(text, quote=False)


def link(url, label):
    return f'<a href="{html.escape(url)}" target="_blank" rel="noopener noreferrer">{esc(label)}</a>'


# ---------- block renderers (each returns list of lines at given indent) ----------

def r_paras(num, start, end, pad):
    return [f"{pad}<p>{esc(t)}</p>" for _, t in lines(num)[start:end]]


def r_joined(num, start, end, pad):
    """One paragraph from source lines that split a single sentence across lines."""
    return [f"{pad}<p>{esc(' '.join(t for _, t in lines(num)[start:end]))}</p>"]


def r_list(items, pad):
    """Nested <ul> from (depth, text) pairs."""
    base = min(d for d, _ in items)
    out, prev = [], -1
    for depth, text in items:
        level = max(0, min(depth - base, prev + 1))
        if level > prev:
            out.append(f"{pad}{'  ' * level}<ul>")
        else:
            if level < prev:
                for closing in range(prev, level, -1):
                    out.append(f"{pad}{'  ' * closing}</li></ul>")
            out.append(f"{pad}{'  ' * level}</li>")
        out.append(f"{pad}{'  ' * level}<li>{esc(text)}")
        prev = level
    for closing in range(prev, -1, -1):
        out.append(f"{pad}{'  ' * closing}</li></ul>")
    return out


STATEMENT_RE = re.compile(r"^(using |public |private |protected |class |void |if\b|if\(|else\b|for\b|for\(|while\b|return\b)")
BRACELESS_RE = re.compile(r"^(if|for|while|else)\b.*[^{;]$|^(if|for|while)\(.*[^{;]$")


def r_code(code_lines, pad, title="Code Example"):
    out = [
        f'{pad}<section class="code-panel code-panel--compact" data-code-type="code">',
        f'{pad}  <div class="code-panel__chrome">',
        f'{pad}    <div class="code-panel__dots">',
        f'{pad}      <span class="code-panel__dot"></span>',
        f'{pad}      <span class="code-panel__dot"></span>',
        f'{pad}      <span class="code-panel__dot"></span>',
        f"{pad}    </div>",
        f'{pad}    <div class="code-panel__title">{esc(title)}</div>',
        f"{pad}  </div>",
        f'{pad}  <div class="code-panel__body" role="region" aria-label="{html.escape(title)}">',
    ]
    level = 0
    pending = False
    for text in code_lines:
        if text.startswith("}"):
            level = max(0, level - 1)
        if text.startswith("//"):
            tone = "comment"
        elif re.fullmatch(r"[{}]+", text):
            tone = "brace"
        elif STATEMENT_RE.match(text):
            tone = "statement"
        else:
            tone = "code"
        shown = min(6, level + (1 if pending and not text.startswith("{") else 0))
        out.append(f'{pad}    <div class="code-panel__line code-panel__line--{tone} indent-level-{shown}">{esc(text)}</div>')
        level = max(0, level + text.count("{") - text.count("}") + (1 if text.startswith("}") else 0))
        pending = bool(BRACELESS_RE.match(text))
    out += [f"{pad}  </div>", f"{pad}</section>"]
    return out


def r_images(imgs, pad):
    out = [f'{pad}<div class="accordion-images--horizontal" style="--accordion-images--horizontal-width: 100%;">']
    for num, idx, width, alt in imgs:
        src = html.escape(images[num][idx]["path"], quote=True)
        out.append(
            f'{pad}  <img src="{src}" class="accordion-image" alt="{html.escape(alt, quote=True)}"'
            f' style="--accordion-image-width: {width};" />'
        )
    out.append(f"{pad}</div>")
    return out


def r_video(video_id, summary, watch_url, pad, start=None):
    s = esc(summary)
    query = f"?start={start}" if start else ""
    return [
        f'{pad}<details class="accordion accordion--nested embed-accordion">',
        f"{pad}  <summary>Video: {s}</summary>",
        f'{pad}  <div class="accordion-body">',
        f'{pad}    <article class="embed-card">',
        f'{pad}      <iframe class="video-embed" src="https://www.youtube.com/embed/{video_id}{query}" title="{html.escape(summary)} video"'
        ' loading="lazy" referrerpolicy="strict-origin-when-cross-origin" allowfullscreen></iframe>',
        f'{pad}      <p><a href="{html.escape(watch_url)}" target="_blank" rel="noopener noreferrer">Open source link</a></p>',
        f"{pad}    </article>",
        f"{pad}  </div>",
        f"{pad}</details>",
    ]


def r_table(num, pad):
    table = next(sh.table for sh in SLIDES[num].shapes if getattr(sh, "has_table", False) and sh.has_table)
    rows = [[re.sub(r"\s+", " ", c.text).strip() for c in row.cells] for row in table.rows]
    out = [f'{pad}<div class="assessment-table-wrap">', f'{pad}  <table class="assessment-table">', f"{pad}    <thead>", f"{pad}      <tr>"]
    out += [f"{pad}        <th>{esc(c)}</th>" for c in rows[0]]
    out += [f"{pad}      </tr>", f"{pad}    </thead>", f"{pad}    <tbody>"]
    for row in rows[1:]:
        out.append(f"{pad}      <tr>")
        out += [f"{pad}        <td>{esc(c)}</td>" for c in row]
        out.append(f"{pad}      </tr>")
    out += [f"{pad}    </tbody>", f"{pad}  </table>", f"{pad}</div>"]
    return out


def r_nested(summary, inner, pad):
    return [
        f'{pad}<details class="accordion accordion--nested">',
        f"{pad}  <summary>{esc(summary)}</summary>",
        f'{pad}  <div class="accordion-body">',
        *inner,
        f"{pad}  </div>",
        f"{pad}</details>",
    ]


def render_block(block, pad):
    kind = block[0]
    if kind == "p":                      # ("p", slide, start, end)
        return r_paras(block[1], block[2], block[3], pad)
    if kind == "pj":                     # ("pj", slide, start, end) -> one joined paragraph
        return r_joined(block[1], block[2], block[3], pad)
    if kind == "text":                   # ("text", literal text)
        return [f"{pad}<p>{esc(block[1])}</p>"]
    if kind == "label":                  # ("label", slide, index) -> h3 from source line
        return [f"{pad}<h3>{esc(lines(block[1])[block[2]][1])}</h3>"]
    if kind == "tree":                   # ("tree", slide, start, end) -> source paragraph levels
        return r_list(lines(block[1])[block[2]:block[3]], pad)
    if kind == "treed":                  # ("treed", slide, [(index, depth), ...]) -> explicit depths
        return r_list([(d, lines(block[1])[i][1]) for i, d in block[2]], pad)
    if kind == "code":                   # ("code", slide, start, end, title[, prefix lines])
        prefix = list(block[5]) if len(block) > 5 else []
        return r_code(prefix + [t for _, t in lines(block[1])[block[2]:block[3]]], pad, block[4])
    if kind == "img":                    # ("img", [(slide, idx, width, alt), ...])
        return r_images(block[1], pad)
    if kind == "video":                  # ("video", id, summary, watch_url[, start])
        return r_video(block[1], block[2], block[3], pad, *(block[4:5] or [None]))
    if kind == "table":                  # ("table", slide)
        return r_table(block[1], pad)
    if kind == "link":                   # ("link", label, url)
        return [f"{pad}<p>{link(block[2], block[1])}</p>"]
    if kind == "refs":                   # ("refs", summary, [(label, url), ...])
        items = [f"{pad}    <ul>"] + [f"{pad}      <li>{link(u, l)}</li>" for l, u in block[2]] + [f"{pad}    </ul>"]
        return r_nested(block[1], items, pad)
    if kind == "tut":                    # ("tut", page, label)
        return [f'{pad}<p><a href="../index.html?page=pages/{block[1]}" target="_top">{esc(block[2])}</a></p>']
    if kind == "from":                   # ("from", slide, index) -> "From <url>" with url isolated for linkify
        m = re.fullmatch(r"(.*?<)(https?://\S+?)(>)", lines(block[1])[block[2]][1])
        return [f"{pad}<p>{esc(m.group(1))}<span>{esc(m.group(2))}</span>{esc(m.group(3))}</p>"]
    if kind == "notes":                  # ("notes", slide[, start, end]) -> speaker notes accordion
        paras = [re.sub(r"\s+", " ", t).strip() for t in SLIDES[block[1]].notes_slide.notes_text_frame.text.split("\n")]
        paras = [t for t in paras if t][(block[2] if len(block) > 2 else None):(block[3] if len(block) > 3 else None)]
        return r_nested("Speaker notes", [f"{pad}    <p>{esc(t)}</p>" for t in paras], pad)
    if kind == "notecode":               # ("notecode", slide, title) -> speaker-notes code as a code panel
        code = [re.sub(r"\s+", " ", t).strip() for t in SLIDES[block[1]].notes_slide.notes_text_frame.text.split("\n")]
        return r_code([t for t in code if t], pad, block[2])
    if kind == "codelist":               # ("codelist", [code lines], title)
        return r_code(block[1], pad, block[2])
    if kind == "dsteps":                 # ("dsteps", slide, start, end) -> "- " lines nested as bullets
        return r_list(dash_steps(block[1], block[2], block[3]), pad)
    if kind == "nested":                 # ("nested", summary, [blocks])
        inner = []
        for b in block[2]:
            inner.extend(render_block(b, pad + "    "))
        return r_nested(block[1], inner, pad)
    raise ValueError(kind)


def heading(num, idx=0):
    return lines(num)[idx][1]




def dash_steps(num, start, end):
    """Source lines prefixed "- " become bullets under the preceding step; "- ...:" lines open a further level."""
    items, nest = [], False
    for _, text in lines(num)[start:end]:
        if text.startswith("- "):
            body = text[2:]
            if body.endswith(":"):
                items.append((1, body))
                nest = True
            else:
                items.append((2 if nest else 1, body))
        else:
            items.append((0, text))
            nest = False
    return items




YT = "https://www.youtube.com/watch?v="
LIGHTING_OVERVIEW = "https://unity3d.com/learn/tutorials/topics/graphics/unity-5-lighting-overview?playlist=17102"


def img(num, idx, width, alt):
    return (num, idx, width, alt)


DUP = "Slide {} in the original deck repeats this content exactly."

# ---------- consolidation plan: 68 source slides -> 15 themed slides ----------
# section: dict(s=summary, src=[source slides], b=[blocks])
PLAN = [
    {"h2": "5MMCS001W Lecture 12 - Unity Lighting and Shaders", "src": "1-4", "sections": [
        {"s": "Lecture Overview", "src": [1, 2], "b": [
            ("p", 1, 0, 3),
            ("p", 2, 1, 2),
        ]},
        {"s": heading(3, 2).rstrip(":"), "src": [3], "b": [
            ("p", 3, 2, 3),
            ("tree", 3, 3, 6),
        ]},
        {"s": heading(4, 1), "src": [4], "b": [
            ("tree", 4, 2, 5),
        ]},
        {"s": "Related Tutorials", "src": [], "b": [
            ("tut", "tut-12-urp-lighting-2026.html", "Tutorial 12: URP Lighting with Caf\u00e9 Model"),
            ("tut", "tut-11-01-blender-materials-and-texturing.html", "Tutorial 11-01: Blender Materials and Texturing"),
        ]},
    ]},
    {"h2": heading(5, 1), "src": "5-7", "sections": [
        {"s": "Features of Direct Lighting", "src": [5], "b": [
            ("p", 5, 2, 3),
            ("tree", 5, 3, 9),
        ]},
        {"s": "Light Component Types", "src": [6], "b": [
            ("tree", 6, 3, 8),
        ]},
        {"s": "Other Light Sources", "src": [7], "b": [
            ("tree", 7, 3, 7),
        ]},
    ]},
    {"h2": "Precomputed and Baked Global Illumination", "src": "8-12", "sections": [
        {"s": heading(8, 1), "src": [8], "b": [
            ("p", 8, 2, 3),
            ("tree", 8, 3, 6),
        ]},
        {"s": heading(9, 1), "src": [9], "b": [
            ("p", 9, 2, 4),
            ("tree", 9, 4, 10),
            ("p", 9, 10, 11),
        ]},
        {"s": heading(10, 1), "src": [10], "b": [
            ("p", 10, 2, 6),
        ]},
        {"s": "Geomerics Enlighten System", "src": [11], "b": [
            ("p", 11, 1, 3),
            ("img", [img(11, 0, "80%", "Enlighten ON versus Enlighten OFF comparison of the same interior scene")]),
        ]},
        {"s": "Precomputed Realtime GI - Features", "src": [12], "b": [
            ("tree", 12, 2, 9),
        ]},
    ]},
    {"h2": "Global Illumination Concepts and Lighting Setup", "src": "13-15", "sections": [
        {"s": heading(13, 1), "src": [13], "b": [
            ("p", 13, 2, 3),
            ("tree", 13, 3, 10),
        ]},
        {"s": heading(14, 1), "src": [14], "b": [
            ("p", 14, 2, 4),
            ("tree", 14, 4, 11),
        ]},
        {"s": heading(15, 1), "src": [15], "b": [
            ("p", 15, 2, 4),
            ("tree", 15, 4, 8),
            ("p", 15, 8, 10),
        ]},
    ]},
    {"h2": "Lighting Panel and Colour Space", "src": "16-19", "sections": [
        {"s": heading(16, 1), "src": [16], "b": [
            ("p", 16, 2, 6),
            ("img", [img(16, 0, "40%", "Unity Lighting panel showing Environment Lighting, Precomputed Realtime GI and Baked GI settings")]),
        ]},
        {"s": heading(17, 1), "src": [17], "b": [
            ("pj", 17, 2, 4),
            ("p", 17, 4, 9),
            ("img", [img(17, 0, "40%", "Unity Lighting panel with Realtime and Baked GI options and the Auto / Build controls")]),
        ]},
        {"s": heading(18, 1), "src": [18], "b": [
            ("p", 18, 2, 3),
            ("tree", 18, 3, 10),
            ("img", [img(18, 0, "45%", "Light component Inspector with the Baking dropdown showing Realtime, Baked and Mixed")]),
        ]},
        {"s": heading(19, 1), "src": [19], "b": [
            ("p", 19, 2, 7),
        ]},
    ]},
    {"h2": "Ambient Light and Skybox", "src": "20-21", "sections": [
        {"s": heading(20, 1), "src": [20], "b": [
            ("p", 20, 2, 6),
            ("tree", 20, 6, 13),
            ("p", 20, 13, 14),
        ]},
        {"s": heading(21, 1), "src": [21], "b": [
            ("video", "vJ6IS2iIiQg", "Unity 5 Intro to Skybox and Realtime Global Illumination", YT + "vJ6IS2iIiQg"),
            ("p", 21, 3, 11),
            ("img", [img(21, 0, "70%", "Skybox and realtime GI tutorial video frame showing a lit desert scene in Unity")]),
        ]},
    ]},
    {"h2": "Light Sources, Render Mode and Light Effects", "src": "22-24", "sections": [
        {"s": heading(22, 1), "src": [22], "b": [
            ("tree", 22, 2, 18),
        ]},
        {"s": heading(23, 1), "src": [23], "b": [
            ("tree", 23, 2, 13),
        ]},
        {"s": heading(24, 1), "src": [24], "b": [
            ("tree", 24, 2, 17),
        ]},
    ]},
    {"h2": "Emissive Materials and Probes", "src": "25-27", "sections": [
        {"s": heading(25, 1), "src": [25, 26], "b": [
            ("p", 25, 2, 6),
            ("p", 26, 2, 3),
        ]},
        {"s": heading(27, 1), "src": [27], "b": [
            ("p", 27, 2, 3),
            ("tree", 27, 3, 8),
        ]},
    ]},
    {"h2": "Unity Standard Shader - Physically Based Rendering", "src": "28-32", "sections": [
        {"s": heading(29, 1), "src": [28, 29], "b": [
            ("video", "fD_ho_ofY6A", "Unity 5 Graphics - The Standard Shader - Unity Official Tutorials", YT + "fD_ho_ofY6A"),
            ("p", 29, 3, 8),
            ("tree", 29, 8, 10),
            ("p", 29, 10, 11),
        ]},
        {"s": "Standard (Metallic) - Metallic and Smoothness", "src": [30], "b": [
            ("p", 30, 2, 4),
            ("tree", 30, 4, 12),
        ]},
        {"s": "Two Choices: Metallic or Specular", "src": [31], "b": [
            ("p", 31, 2, 4),
            ("tree", 31, 4, 6),
            ("p", 31, 6, 7),
            ("img", [img(31, 0, "90%", "Spheres showing metalness and smoothness values from 0 to 1")]),
        ]},
        {"s": "Smoothness in Metallic and Specular", "src": [32], "b": [
            ("p", 32, 1, 5),
        ]},
    ]},
    {"h2": "Metallic vs Specular Workflows and Calibration", "src": "33-36", "sections": [
        {"s": "Metallic Workflow", "src": [33], "b": [
            ("p", 33, 1, 4),
        ]},
        {"s": "Specular Workflow", "src": [34], "b": [
            ("p", 34, 1, 5),
        ]},
        {"s": "Physically Based Shading Values - Material Charts", "src": [35, 36], "b": [
            ("p", 35, 2, 6),
            ("img", [img(36, 0, "45%", "Unity Shader Calibration Scene chart: Metallic value charts for albedo, metallic and smoothness"),
                     img(36, 1, "45%", "Unity Shader Calibration Scene chart: Specular value charts for albedo, specular and smoothness")]),
        ]},
    ]},
    {"h2": "Standard Shader Maps, Rendering Modes and Materials", "src": "37-41", "sections": [
        {"s": "Standard Shader - Maps (Both have...)", "src": [37], "b": [
            ("p", 37, 2, 4),
            ("tree", 37, 4, 18),
        ]},
        {"s": "Standard Shader - Rendering Mode", "src": [38], "b": [
            ("p", 38, 2, 5),
            ("tree", 38, 5, 16),
        ]},
        {"s": "Standard Shader - Specular Smoothness", "src": [39], "b": [
            ("p", 39, 2, 5),
            ("tree", 39, 5, 10),
        ]},
        {"s": heading(40, 1), "src": [40, 41], "b": [
            ("p", 40, 2, 3),
            ("tree", 40, 3, 7),
            ("text", "Slide 41 in the original deck repeats these steps with the diagram below."),
            ("img", [img(40, 0, "40%", "Inspector for a Capsule showing Mesh Renderer material slots, light and reflection probe settings"),
                     img(41, 0, "45%", "Diagram: 3D Model (vertex positions, colours, normals, UV data), Material (textures, shader property values) and Shader (CG / HLSL code)")]),
        ]},
    ]},
    {"h2": "Lighting and Shader Video Tutorials", "src": "42-45", "sections": [
        {"s": heading(42, 1), "src": [42], "b": [
            ("text", "Video embedded under Unity Standard Shader - Physically Based Rendering."),
            ("img", [img(42, 0, "70%", "Standard Shader tutorial video frame with a crate and fridge model in Unity")]),
        ]},
        {"s": heading(43, 1), "src": [43], "b": [
            ("text", "Video embedded under Ambient Light and Skybox."),
            ("img", [img(43, 0, "70%", "Skybox GI tutorial video frame showing the Unity editor with a lit scene")]),
        ]},
        {"s": heading(44, 1), "src": [44], "b": [
            ("p", 44, 2, 5),
            ("text", "Standard Shader video: embedded under Unity Standard Shader - Physically Based Rendering."),
        ]},
        {"s": heading(45, 1), "src": [45], "b": [
            ("video", "Ukkiqe_75d0", "Unity 5 - 100% Physically based shading - demo scene - full day/night", YT + "Ukkiqe_75d0"),
            ("video", "h5GFoI38DOg", "Unity 5 Tutorial - Realtime Global Illumination, Day/Night Cycle, Reflection Probes", YT + "h5GFoI38DOg"),
            ("img", [img(45, 0, "60%", "Physically based shading demo: sunlit corridor opening onto the sea")]),
        ]},
    ]},
    {"h2": "URP Updates, Conversions - Quick Steps", "src": "46-50", "sections": [
        {"s": heading(46, 2), "src": [46], "b": [
            ("tree", 46, 3, 11),
        ]},
        {"s": "Steps 1-3: Backup, Install URP, Check URP Asset", "src": [47], "b": [
            ("dsteps", 47, 2, 10),
        ]},
        {"s": "Steps 4-5: Create Pipeline Asset, Quality Levels", "src": [48], "b": [
            ("dsteps", 48, 2, 8),
        ]},
        {"s": "Steps 6-7: Upgrade Materials, Re-bake Lighting", "src": [49], "b": [
            ("dsteps", 49, 2, 8),
        ]},
        {"s": "Steps 8-9: Quick Check and Fixes", "src": [50], "b": [
            ("dsteps", 50, 2, 9),
        ]},
    ]},
    {"h2": "URP Conversion in Detail - Setup and Materials", "src": "51-57", "sections": [
        {"s": heading(51, 2) + " (detailed)", "src": [51], "b": [
            ("tree", 51, 3, 11),
        ]},
        {"s": heading(52, 2), "src": [52], "b": [
            ("dsteps", 52, 3, 9),
        ]},
        {"s": heading(53, 2), "src": [53], "b": [
            ("dsteps", 53, 3, 8),
        ]},
        {"s": heading(54, 2), "src": [54, 55], "b": [
            ("dsteps", 54, 3, 7),
            ("dsteps", 55, 2, 6),
        ]},
        {"s": heading(56, 2), "src": [56, 57], "b": [
            ("dsteps", 56, 3, 10),
            ("dsteps", 57, 2, 6),
        ]},
    ]},
    {"h2": "URP Conversion in Detail - Shaders, Renderer, Lights, Shadows", "src": "58-68", "sections": [
        {"s": heading(58, 2), "src": [58, 59], "b": [
            ("dsteps", 58, 3, 9),
            ("dsteps", 59, 3, 8),
        ]},
        {"s": heading(60, 2) + " - Pipeline Asset", "src": [60, 61, 62], "b": [
            ("dsteps", 60, 3, 7),
            ("dsteps", 61, 3, 14),
            ("text", DUP.format(62)),
        ]},
        {"s": heading(60, 2) + " - Forward Renderer", "src": [63], "b": [
            ("dsteps", 63, 2, 11),
        ]},
        {"s": heading(64, 2), "src": [64, 65, 66], "b": [
            ("dsteps", 64, 3, 9),
            ("dsteps", 65, 3, 11),
            ("text", DUP.format(66)),
        ]},
        {"s": heading(67, 2), "src": [67, 68], "b": [
            ("dsteps", 67, 3, 10),
            ("text", DUP.format(68)),
        ]},
    ]},
]


def render_section(sec, number, sub, first):
    pad = "          "
    body = []
    for block in sec["b"]:
        body.extend(render_block(block, pad))
    src = ", ".join(str(n) for n in sec["src"]) or "added cross-reference"
    return [
        f"      <!-- Section {number}.{sub}: {sec['s']} (source slide {src}) -->",
        f'      <details class="accordion"{" open" if first else ""}>',
        f"        <summary>{esc(sec['s'])}</summary>",
        '        <div class="accordion-body">',
        *body,
        "        </div>",
        "      </details>",
    ]


def render_slide(spec, number, total):
    first = number == 1
    out = [
        f"    <!-- Slide {number}: {spec['h2']} (source slides {spec['src']}) -->",
        f'    <div class="slide{" active" if first else ""}">',
        f'      <div class="slide-title">{SLIDE_TITLE}</div>',
        '      <div class="slide-content slide-content-accordion">',
        f"        <h2>{esc(spec['h2'])}</h2>",
        '        <div class="accordion-scroll">',
    ]
    for i, sec in enumerate(spec["sections"]):
        out.extend(render_section(sec, number, i + 1, i == 0))
    out += [
        "        </div>",
        "      </div>",
        '      <div class="slide-controls">',
        f'        <button class="slide-button" data-slide-step="-1"{" disabled" if first else ""}> ← Previous</button>',
        '        <img src="images/UOW_Logo_Length_Alpha.png" class="slide-logo" alt="University of Westminster" />',
        f'        <button class="slide-button" data-slide-step="1"{" disabled" if number == total else ""}>Next → </button>',
        f'        <div class="slide-counter"><span id="current-slide-{number}">{number}</span>'
        f' / <span id="total-slides">{total}</span></div>',
        "      </div>",
        '      <div class="keyboard-hint">',
        "        Use <kbd>→</kbd> <kbd>←</kbd> or click buttons to navigate",
        "      </div>",
        "    </div>",
        "",
    ]
    return out


def main():
    covered = {n for spec in PLAN for sec in spec["sections"] for n in sec["src"]}
    missing = set(SLIDES) - covered
    if missing:
        raise SystemExit(f"Source slides not mapped: {sorted(missing)}")

    original = FOOTER_SOURCE.read_text(encoding="utf-8").split("\n")
    marker = '  <script src="../assets/js/site.js'
    footer = original[next(i for i, ln in enumerate(original) if ln.startswith(marker)):]

    total = len(PLAN)
    doc = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '  <meta charset="UTF-8" />',
        '  <meta name="viewport" content="width=device-width, initial-scale=1.0" />',
        f"  <title>{PAGE_TITLE}</title>",
        '  <link rel="preconnect" href="https://fonts.googleapis.com" />',
        '  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />',
        '  <link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;600;700&family=Space+Grotesk:wght@500;700&display=swap" rel="stylesheet" />',
        '  <link rel="stylesheet" href="../assets/css/site.css?v=20261007a" />',
        '  <link rel="stylesheet" href="../assets/css/slideshow.css?v=20261007a" />',
        "</head>",
        '<body class="lecture-fullbleed">',
        '  <div class="slideshow-container">',
    ]
    for i, spec in enumerate(PLAN, 1):
        doc.extend(render_slide(spec, i, total))
    doc += ["  </div>", ""] + footer

    OUT_PATH.write_text("\n".join(doc), encoding="utf-8")
    print(f"wrote {OUT_PATH} with {total} slides from {len(SLIDES)} source slides")


if __name__ == "__main__":
    main()
