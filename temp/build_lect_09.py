"""Build lect-09-rb-collision.html with consolidated (theme-grouped) slides.

Text, paragraph levels, tables, hyperlinks and speaker notes are read from the
PPTX directly (the extractor JSON flattens levels and drops tables); images use
the extractor output in lect-09-rb-collision-content.json. The PLAN
maps every source slide (1-23) onto a smaller set of themed slides; see
docs/lecture-conversion-playbook.md.
"""
import html
import json
import re
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PAGES = ROOT / "site" / "pages"
PPTX_PATH = ROOT / "site" / "source-lecture-pptx" / "3DIMD Lecture 09 - XR Week 10 - UNITY RB and Collision 04.pptx"
JSON_PATH = PAGES / "lect-09-rb-collision-content.json"
OUT_PATH = PAGES / "lect-09-rb-collision.html"
FOOTER_SOURCE = PAGES / "lect-04a-report.html"

SLIDE_TITLE = "3D Interactive Media Development"
PAGE_TITLE = "3DIMD | Lecture 09 - Unity Rigidbodies, Physics and Collision"

images = {s["slide_number"]: s["images"] for s in json.load(JSON_PATH.open(encoding="utf-8"))}
prs = Presentation(str(PPTX_PATH))
SLIDES = {i: s for i, s in enumerate(prs.slides, 1)}


def _collect_lines(slide):
    out = []

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6:
                walk(shape.shapes)
                continue
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for seg in para.text.replace("\x0b", "\n").replace("\r", "\n").split("\n"):
                    lead = re.match(r"^[\t\u2003\xa0 ]*", seg).group(0)
                    text = re.sub(r"\s+", " ", seg.replace("\xa0", " ")).strip()
                    if text and text != SLIDE_TITLE:
                        out.append((para.level + lead.count("\t"), text))

    walk(slide.shapes)
    return out


LINES = {n: _collect_lines(s) for n, s in SLIDES.items()}


def lines(num):
    return LINES[num]


def esc(text):
    return html.escape(text, quote=False)


def link(url, label):
    return f'<a href="{html.escape(url)}" target="_blank" rel="noopener noreferrer">{esc(label)}</a>'


# ---------- block renderers (each returns list of lines at given indent) ----------

def r_paras(num, start, end, pad):
    return [f"{pad}<p>{esc(t)}</p>" for _, t in lines(num)[start:end]]


def r_joined(num, start, end, pad):
    """One paragraph from source lines that split a single sentence across lines."""
    return [f"{pad}<p>{esc(' '.join(t for _, t in lines(num)[start:end]))}</p>"]


def r_list(items, pad):
    """Nested <ul> from (depth, text) pairs."""
    base = min(d for d, _ in items)
    out, prev = [], -1
    for depth, text in items:
        level = max(0, min(depth - base, prev + 1))
        if level > prev:
            out.append(f"{pad}{'  ' * level}<ul>")
        else:
            if level < prev:
                for closing in range(prev, level, -1):
                    out.append(f"{pad}{'  ' * closing}</li></ul>")
            out.append(f"{pad}{'  ' * level}</li>")
        out.append(f"{pad}{'  ' * level}<li>{esc(text)}")
        prev = level
    for closing in range(prev, -1, -1):
        out.append(f"{pad}{'  ' * closing}</li></ul>")
    return out


STATEMENT_RE = re.compile(r"^(using |public |private |protected |class |void |if\b|if\(|else\b|for\b|for\(|while\b|return\b)")
BRACELESS_RE = re.compile(r"^(if|for|while|else)\b.*[^{;]$|^(if|for|while)\(.*[^{;]$")


def r_code(code_lines, pad, title="Code Example"):
    out = [
        f'{pad}<section class="code-panel code-panel--compact" data-code-type="code">',
        f'{pad}  <div class="code-panel__chrome">',
        f'{pad}    <div class="code-panel__dots">',
        f'{pad}      <span class="code-panel__dot"></span>',
        f'{pad}      <span class="code-panel__dot"></span>',
        f'{pad}      <span class="code-panel__dot"></span>',
        f"{pad}    </div>",
        f'{pad}    <div class="code-panel__title">{esc(title)}</div>',
        f"{pad}  </div>",
        f'{pad}  <div class="code-panel__body" role="region" aria-label="{html.escape(title)}">',
    ]
    level = 0
    pending = False
    for text in code_lines:
        if text.startswith("}"):
            level = max(0, level - 1)
        if text.startswith("//"):
            tone = "comment"
        elif re.fullmatch(r"[{}]+", text):
            tone = "brace"
        elif STATEMENT_RE.match(text):
            tone = "statement"
        else:
            tone = "code"
        shown = min(6, level + (1 if pending and not text.startswith("{") else 0))
        out.append(f'{pad}    <div class="code-panel__line code-panel__line--{tone} indent-level-{shown}">{esc(text)}</div>')
        if text.endswith("{"):
            level += 1
        pending = bool(BRACELESS_RE.match(text))
    out += [f"{pad}  </div>", f"{pad}</section>"]
    return out


def r_images(imgs, pad):
    out = [f'{pad}<div class="accordion-images--horizontal" style="--accordion-images--horizontal-width: 100%;">']
    for num, idx, width, alt in imgs:
        src = html.escape(images[num][idx]["path"], quote=True)
        out.append(
            f'{pad}  <img src="{src}" class="accordion-image" alt="{html.escape(alt, quote=True)}"'
            f' style="--accordion-image-width: {width};" />'
        )
    out.append(f"{pad}</div>")
    return out


def r_video(video_id, summary, watch_url, pad, start=None):
    s = esc(summary)
    query = f"?start={start}" if start else ""
    return [
        f'{pad}<details class="accordion accordion--nested embed-accordion">',
        f"{pad}  <summary>Video: {s}</summary>",
        f'{pad}  <div class="accordion-body">',
        f'{pad}    <article class="embed-card">',
        f'{pad}      <iframe class="video-embed" src="https://www.youtube.com/embed/{video_id}{query}" title="{html.escape(summary)} video"'
        ' loading="lazy" referrerpolicy="strict-origin-when-cross-origin" allowfullscreen></iframe>',
        f'{pad}      <p><a href="{html.escape(watch_url)}" target="_blank" rel="noopener noreferrer">Open source link</a></p>',
        f"{pad}    </article>",
        f"{pad}  </div>",
        f"{pad}</details>",
    ]


def r_table(num, pad):
    table = next(sh.table for sh in SLIDES[num].shapes if getattr(sh, "has_table", False) and sh.has_table)
    rows = [[re.sub(r"\s+", " ", c.text).strip() for c in row.cells] for row in table.rows]
    out = [f'{pad}<div class="assessment-table-wrap">', f'{pad}  <table class="assessment-table">', f"{pad}    <thead>", f"{pad}      <tr>"]
    out += [f"{pad}        <th>{esc(c)}</th>" for c in rows[0]]
    out += [f"{pad}      </tr>", f"{pad}    </thead>", f"{pad}    <tbody>"]
    for row in rows[1:]:
        out.append(f"{pad}      <tr>")
        out += [f"{pad}        <td>{esc(c)}</td>" for c in row]
        out.append(f"{pad}      </tr>")
    out += [f"{pad}    </tbody>", f"{pad}  </table>", f"{pad}</div>"]
    return out


def r_nested(summary, inner, pad):
    return [
        f'{pad}<details class="accordion accordion--nested">',
        f"{pad}  <summary>{esc(summary)}</summary>",
        f'{pad}  <div class="accordion-body">',
        *inner,
        f"{pad}  </div>",
        f"{pad}</details>",
    ]


def render_block(block, pad):
    kind = block[0]
    if kind == "p":                      # ("p", slide, start, end)
        return r_paras(block[1], block[2], block[3], pad)
    if kind == "pj":                     # ("pj", slide, start, end) -> one joined paragraph
        return r_joined(block[1], block[2], block[3], pad)
    if kind == "text":                   # ("text", literal text)
        return [f"{pad}<p>{esc(block[1])}</p>"]
    if kind == "label":                  # ("label", slide, index) -> h3 from source line
        return [f"{pad}<h3>{esc(lines(block[1])[block[2]][1])}</h3>"]
    if kind == "tree":                   # ("tree", slide, start, end) -> source paragraph levels
        return r_list(lines(block[1])[block[2]:block[3]], pad)
    if kind == "treed":                  # ("treed", slide, [(index, depth), ...]) -> explicit depths
        return r_list([(d, lines(block[1])[i][1]) for i, d in block[2]], pad)
    if kind == "code":                   # ("code", slide, start, end, title[, prefix lines])
        prefix = list(block[5]) if len(block) > 5 else []
        return r_code(prefix + [t for _, t in lines(block[1])[block[2]:block[3]]], pad, block[4])
    if kind == "img":                    # ("img", [(slide, idx, width, alt), ...])
        return r_images(block[1], pad)
    if kind == "video":                  # ("video", id, summary, watch_url[, start])
        return r_video(block[1], block[2], block[3], pad, *(block[4:5] or [None]))
    if kind == "table":                  # ("table", slide)
        return r_table(block[1], pad)
    if kind == "link":                   # ("link", label, url)
        return [f"{pad}<p>{link(block[2], block[1])}</p>"]
    if kind == "refs":                   # ("refs", summary, [(label, url), ...])
        items = [f"{pad}    <ul>"] + [f"{pad}      <li>{link(u, l)}</li>" for l, u in block[2]] + [f"{pad}    </ul>"]
        return r_nested(block[1], items, pad)
    if kind == "tut":                    # ("tut", page, label)
        return [f'{pad}<p><a href="../index.html?page=pages/{block[1]}" target="_top">{esc(block[2])}</a></p>']
    if kind == "from":                   # ("from", slide, index) -> "From <url>" with url isolated for linkify
        m = re.fullmatch(r"(.*?<)(https?://\S+?)(>)", lines(block[1])[block[2]][1])
        return [f"{pad}<p>{esc(m.group(1))}<span>{esc(m.group(2))}</span>{esc(m.group(3))}</p>"]
    if kind == "nested":                 # ("nested", summary, [blocks])
        inner = []
        for b in block[2]:
            inner.extend(render_block(b, pad + "    "))
        return r_nested(block[1], inner, pad)
    raise ValueError(kind)


def heading(num, idx=0):
    return lines(num)[idx][1]


YT = "https://www.youtube.com/watch?v="

# ---------- consolidation plan: 23 source slides -> 8 themed slides ----------
# Line 0 of every source slide is the repeated deck header ("3DIMD and XR Multimodal
# Interaction" / "XR Multimodal Interaction"); it is shown once in the overview.
# section: dict(s=summary, src=[source slides], b=[blocks])
PLAN = [
    {"h2": "Lecture 09 - Unity Physics / Collision", "src": "1-2", "sections": [
        {"s": "Lecture Overview", "src": [1, 2], "b": [
            ("p", 1, 0, 3),
            ("p", 2, 1, 2),
        ]},
        {"s": "Interactive Examples", "src": [1], "b": [
            ("p", 1, 3, 4),
            ("video", "gwYjnWCcw18", "Immerse Framework for Unity3D", YT + "gwYjnWCcw18"),
        ]},
        {"s": "Related Tutorials", "src": [], "b": [
            ("tut", "tut-02-03-students.html", "Tutorials 02–03: Input, movement and collision exercises"),
            ("tut", "tut-08-physics-quiz-rbct.html", "Tutorial 08 Quiz: Rigidbodies, colliders and triggers"),
            ("tut", "tut-08-character-troubleshooting-bb.html", "Tutorial 08: Character controllers and troubleshooting"),
        ]},
    ]},
    {"h2": "GameObject Dynamics Options - Static, Kinematic and Dynamic", "src": "3-4, 14", "sections": [
        {"s": heading(3, 1), "src": [3], "b": [
            ("video", "xp37Hz1t1Q8", "Unity Physics: Static, Kinematic, Dynamic", YT + "xp37Hz1t1Q8"),
        ]},
        {"s": "Static, Kinematic RigidBody and Dynamic RigidBody", "src": [4], "b": [
            ("tree", 4, 1, 10),
        ]},
        {"s": heading(14, 1), "src": [14], "b": [
            ("label", 14, 2),
            ("tree", 14, 3, 15),
        ]},
    ]},
    {"h2": "Unity Collisions and Triggers", "src": "5-8", "sections": [
        {"s": heading(6, 1), "src": [5, 6], "b": [
            ("tree", 6, 2, 10),
        ]},
        {"s": heading(7, 1), "src": [7, 8], "b": [
            ("img", [(8, 0, "100%", "Unity collider and trigger matrix: which collider combinations produce collision or trigger messages")]),
        ]},
    ]},
    {"h2": "Unity Collisions for Triggers", "src": "9-10", "sections": [
        {"s": heading(9, 1), "src": [9, 10], "b": [
            ("p", 10, 1, 2),
            ("tree", 10, 2, 14),
        ]},
    ]},
    {"h2": "Unity Collisions - True Collision Detection", "src": "11-13, 18", "sections": [
        {"s": heading(11, 1), "src": [11, 12], "b": [
            ("label", 12, 1),
            ("p", 12, 2, 4),
            ("tree", 12, 4, 9),
        ]},
        {"s": "Gravity, Static, IsKinematic and Dynamic", "src": [13], "b": [
            ("p", 13, 2, 3),
            ("treed", 13, [(4, 0), (5, 0), (6, 0), (7, 1), (8, 0), (9, 1), (10, 1), (11, 0)]),
            ("video", "S0MXZ6AXzTE", "Unity Explained - Colliders - Beginner Tutorial", YT + "S0MXZ6AXzTE"),
        ]},
        {"s": heading(18, 4), "src": [18], "b": [
            ("tree", 18, 5, 9),
        ]},
    ]},
    {"h2": "Linking Rigidbodies to the XR Interaction Toolkit (XRIT)", "src": "15-17", "sections": [
        {"s": heading(15, 1), "src": [15, 17], "b": [
            ("p", 15, 2, 4),
            ("p", 17, 2, 5),
        ]},
        {"s": heading(16, 1), "src": [16], "b": [
            ("treed", 16, [(2, 0), (3, 1), (4, 2), (5, 1), (6, 2), (7, 1), (8, 2)]),
        ]},
        {"s": heading(17, 5), "src": [17], "b": [
            ("tree", 17, 6, 9),
            ("p", 17, 9, 10),
        ]},
    ]},
    {"h2": "Unity Collisions and Triggers Test Project", "src": "19-21", "sections": [
        {"s": heading(19, 1), "src": [19], "b": [
            ("video", "WFkbqdo2OI4", "Unity 3d: Collisions, Triggers, Rigidbodies, and Kinematics", YT + "WFkbqdo2OI4"),
        ]},
        {"s": heading(20, 1), "src": [20], "b": [
            ("img", [(20, 0, "70%", "CollisionTester script with OnCollisionEnter, OnTriggerEnter and FixedUpdate movement")]),
        ]},
        {"s": heading(21, 3), "src": [21], "b": [
            ("tree", 21, 4, 14),
        ]},
    ]},
    {"h2": "Character Controllers and Next Steps", "src": "22-23", "sections": [
        {"s": heading(22, 1), "src": [22], "b": [
            ("label", 22, 2),
            ("tree", 22, 5, 8),
            ("video", "e94KggaEAr4", heading(22, 3), YT + "e94KggaEAr4"),
            ("p", 22, 4, 5),
        ]},
        {"s": heading(22, 8), "src": [22], "b": [
            ("video", "qdskE8PJy6Q", heading(22, 9), YT + "qdskE8PJy6Q&t=9s", 9),
            ("p", 22, 10, 11),
        ]},
        {"s": heading(22, 11), "src": [22], "b": [
            ("video", "2MkNoaWJWf4", heading(22, 12), YT + "2MkNoaWJWf4"),
            ("p", 22, 13, 14),
            ("from", 22, 14),
            ("p", 22, 15, 16),
            ("from", 22, 16),
            ("p", 22, 17, 18),
            ("video", "ly9mK0TGJJo", heading(22, 18), YT + "ly9mK0TGJJo"),
            ("p", 22, 19, 21),
        ]},
        {"s": heading(23, 1), "src": [23], "b": [
            ("p", 23, 2, 3),
        ]},
    ]},
]


def render_section(sec, number, sub, first):
    pad = "          "
    body = []
    for block in sec["b"]:
        body.extend(render_block(block, pad))
    src = ", ".join(str(n) for n in sec["src"]) or "added cross-reference"
    return [
        f"      <!-- Section {number}.{sub}: {sec['s']} (source slide {src}) -->",
        f'      <details class="accordion"{" open" if first else ""}>',
        f"        <summary>{esc(sec['s'])}</summary>",
        '        <div class="accordion-body">',
        *body,
        "        </div>",
        "      </details>",
    ]


def render_slide(spec, number, total):
    first = number == 1
    out = [
        f"    <!-- Slide {number}: {spec['h2']} (source slides {spec['src']}) -->",
        f'    <div class="slide{" active" if first else ""}">',
        f'      <div class="slide-title">{SLIDE_TITLE}</div>',
        '      <div class="slide-content slide-content-accordion">',
        f"        <h2>{esc(spec['h2'])}</h2>",
        '        <div class="accordion-scroll">',
    ]
    for i, sec in enumerate(spec["sections"]):
        out.extend(render_section(sec, number, i + 1, i == 0))
    out += [
        "        </div>",
        "      </div>",
        '      <div class="slide-controls">',
        f'        <button class="slide-button" data-slide-step="-1"{" disabled" if first else ""}> ← Previous</button>',
        '        <img src="images/UOW_Logo_Length_Alpha.png" class="slide-logo" alt="University of Westminster" />',
        f'        <button class="slide-button" data-slide-step="1"{" disabled" if number == total else ""}>Next → </button>',
        f'        <div class="slide-counter"><span id="current-slide-{number}">{number}</span>'
        f' / <span id="total-slides">{total}</span></div>',
        "      </div>",
        '      <div class="keyboard-hint">',
        "        Use <kbd>→</kbd> <kbd>←</kbd> or click buttons to navigate",
        "      </div>",
        "    </div>",
        "",
    ]
    return out


def main():
    covered = {n for spec in PLAN for sec in spec["sections"] for n in sec["src"]}
    missing = set(SLIDES) - covered
    if missing:
        raise SystemExit(f"Source slides not mapped: {sorted(missing)}")

    original = FOOTER_SOURCE.read_text(encoding="utf-8").split("\n")
    marker = '  <script src="../assets/js/site.js'
    footer = original[next(i for i, ln in enumerate(original) if ln.startswith(marker)):]

    total = len(PLAN)
    doc = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '  <meta charset="UTF-8" />',
        '  <meta name="viewport" content="width=device-width, initial-scale=1.0" />',
        f"  <title>{PAGE_TITLE}</title>",
        '  <link rel="preconnect" href="https://fonts.googleapis.com" />',
        '  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />',
        '  <link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;600;700&family=Space+Grotesk:wght@500;700&display=swap" rel="stylesheet" />',
        '  <link rel="stylesheet" href="../assets/css/site.css?v=20261007a" />',
        '  <link rel="stylesheet" href="../assets/css/slideshow.css?v=20261007a" />',
        "</head>",
        '<body class="lecture-fullbleed">',
        '  <div class="slideshow-container">',
    ]
    for i, spec in enumerate(PLAN, 1):
        doc.extend(render_slide(spec, i, total))
    doc += ["  </div>", ""] + footer

    OUT_PATH.write_text("\n".join(doc), encoding="utf-8")
    print(f"wrote {OUT_PATH} with {total} slides from {len(SLIDES)} source slides")


if __name__ == "__main__":
    main()
