"""Build lect-11-cameras-lerp.html with consolidated (theme-grouped) slides.

Text, paragraph levels, tables, hyperlinks and speaker notes are read from the
PPTX directly (the extractor JSON flattens levels and drops tables); images use
the extractor output in lect-11-cameras-lerp-content.json. The PLAN
maps every source slide (1-55) onto a smaller set of themed slides; see
docs/lecture-conversion-playbook.md.
"""
import html
import json
import re
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PAGES = ROOT / "site" / "pages"
PPTX_PATH = ROOT / "site" / "source-lecture-pptx" / "3DIMD Lecture 11 Cameras and Lerp.pptx"
JSON_PATH = PAGES / "lect-11-cameras-lerp-content.json"
OUT_PATH = PAGES / "lect-11-cameras-lerp.html"
FOOTER_SOURCE = PAGES / "lect-04a-report.html"

SLIDE_TITLE = "3D Interactive Media Development"
PAGE_TITLE = "3DIMD | Lecture 11 - Cameras and Lerp"

images = {s["slide_number"]: s["images"] for s in json.load(JSON_PATH.open(encoding="utf-8"))}
prs = Presentation(str(PPTX_PATH))
SLIDES = {i: s for i, s in enumerate(prs.slides, 1)}


def _collect_lines(slide):
    out = []

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6:
                walk(shape.shapes)
                continue
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for seg in para.text.replace("\x0b", "\n").replace("\r", "\n").split("\n"):
                    lead = re.match(r"^[\t\u2003\xa0 ]*", seg).group(0)
                    text = re.sub(r"\s+", " ", seg.replace("\xa0", " ")).strip()
                    if text and text != SLIDE_TITLE:
                        out.append((para.level + lead.count("\t"), text))

    walk(slide.shapes)
    return out


LINES = {n: _collect_lines(s) for n, s in SLIDES.items()}


def lines(num):
    return LINES[num]


def esc(text):
    return html.escape(text, quote=False)


def link(url, label):
    return f'<a href="{html.escape(url)}" target="_blank" rel="noopener noreferrer">{esc(label)}</a>'


# ---------- block renderers (each returns list of lines at given indent) ----------

def r_paras(num, start, end, pad):
    return [f"{pad}<p>{esc(t)}</p>" for _, t in lines(num)[start:end]]


def r_joined(num, start, end, pad):
    """One paragraph from source lines that split a single sentence across lines."""
    return [f"{pad}<p>{esc(' '.join(t for _, t in lines(num)[start:end]))}</p>"]


def r_list(items, pad):
    """Nested <ul> from (depth, text) pairs."""
    base = min(d for d, _ in items)
    out, prev = [], -1
    for depth, text in items:
        level = max(0, min(depth - base, prev + 1))
        if level > prev:
            out.append(f"{pad}{'  ' * level}<ul>")
        else:
            if level < prev:
                for closing in range(prev, level, -1):
                    out.append(f"{pad}{'  ' * closing}</li></ul>")
            out.append(f"{pad}{'  ' * level}</li>")
        out.append(f"{pad}{'  ' * level}<li>{esc(text)}")
        prev = level
    for closing in range(prev, -1, -1):
        out.append(f"{pad}{'  ' * closing}</li></ul>")
    return out


STATEMENT_RE = re.compile(r"^(using |public |private |protected |class |void |if\b|if\(|else\b|for\b|for\(|while\b|return\b)")
BRACELESS_RE = re.compile(r"^(if|for|while|else)\b.*[^{;]$|^(if|for|while)\(.*[^{;]$")


def r_code(code_lines, pad, title="Code Example"):
    out = [
        f'{pad}<section class="code-panel code-panel--compact" data-code-type="code">',
        f'{pad}  <div class="code-panel__chrome">',
        f'{pad}    <div class="code-panel__dots">',
        f'{pad}      <span class="code-panel__dot"></span>',
        f'{pad}      <span class="code-panel__dot"></span>',
        f'{pad}      <span class="code-panel__dot"></span>',
        f"{pad}    </div>",
        f'{pad}    <div class="code-panel__title">{esc(title)}</div>',
        f"{pad}  </div>",
        f'{pad}  <div class="code-panel__body" role="region" aria-label="{html.escape(title)}">',
    ]
    level = 0
    pending = False
    for text in code_lines:
        if text.startswith("}"):
            level = max(0, level - 1)
        if text.startswith("//"):
            tone = "comment"
        elif re.fullmatch(r"[{}]+", text):
            tone = "brace"
        elif STATEMENT_RE.match(text):
            tone = "statement"
        else:
            tone = "code"
        shown = min(6, level + (1 if pending and not text.startswith("{") else 0))
        out.append(f'{pad}    <div class="code-panel__line code-panel__line--{tone} indent-level-{shown}">{esc(text)}</div>')
        level = max(0, level + text.count("{") - text.count("}") + (1 if text.startswith("}") else 0))
        pending = bool(BRACELESS_RE.match(text))
    out += [f"{pad}  </div>", f"{pad}</section>"]
    return out


def r_images(imgs, pad):
    out = [f'{pad}<div class="accordion-images--horizontal" style="--accordion-images--horizontal-width: 100%;">']
    for num, idx, width, alt in imgs:
        src = html.escape(images[num][idx]["path"], quote=True)
        out.append(
            f'{pad}  <img src="{src}" class="accordion-image" alt="{html.escape(alt, quote=True)}"'
            f' style="--accordion-image-width: {width};" />'
        )
    out.append(f"{pad}</div>")
    return out


def r_video(video_id, summary, watch_url, pad, start=None):
    s = esc(summary)
    query = f"?start={start}" if start else ""
    return [
        f'{pad}<details class="accordion accordion--nested embed-accordion">',
        f"{pad}  <summary>Video: {s}</summary>",
        f'{pad}  <div class="accordion-body">',
        f'{pad}    <article class="embed-card">',
        f'{pad}      <iframe class="video-embed" src="https://www.youtube.com/embed/{video_id}{query}" title="{html.escape(summary)} video"'
        ' loading="lazy" referrerpolicy="strict-origin-when-cross-origin" allowfullscreen></iframe>',
        f'{pad}      <p><a href="{html.escape(watch_url)}" target="_blank" rel="noopener noreferrer">Open source link</a></p>',
        f"{pad}    </article>",
        f"{pad}  </div>",
        f"{pad}</details>",
    ]


def r_table(num, pad):
    table = next(sh.table for sh in SLIDES[num].shapes if getattr(sh, "has_table", False) and sh.has_table)
    rows = [[re.sub(r"\s+", " ", c.text).strip() for c in row.cells] for row in table.rows]
    out = [f'{pad}<div class="assessment-table-wrap">', f'{pad}  <table class="assessment-table">', f"{pad}    <thead>", f"{pad}      <tr>"]
    out += [f"{pad}        <th>{esc(c)}</th>" for c in rows[0]]
    out += [f"{pad}      </tr>", f"{pad}    </thead>", f"{pad}    <tbody>"]
    for row in rows[1:]:
        out.append(f"{pad}      <tr>")
        out += [f"{pad}        <td>{esc(c)}</td>" for c in row]
        out.append(f"{pad}      </tr>")
    out += [f"{pad}    </tbody>", f"{pad}  </table>", f"{pad}</div>"]
    return out


def r_nested(summary, inner, pad):
    return [
        f'{pad}<details class="accordion accordion--nested">',
        f"{pad}  <summary>{esc(summary)}</summary>",
        f'{pad}  <div class="accordion-body">',
        *inner,
        f"{pad}  </div>",
        f"{pad}</details>",
    ]


def render_block(block, pad):
    kind = block[0]
    if kind == "p":                      # ("p", slide, start, end)
        return r_paras(block[1], block[2], block[3], pad)
    if kind == "pj":                     # ("pj", slide, start, end) -> one joined paragraph
        return r_joined(block[1], block[2], block[3], pad)
    if kind == "text":                   # ("text", literal text)
        return [f"{pad}<p>{esc(block[1])}</p>"]
    if kind == "label":                  # ("label", slide, index) -> h3 from source line
        return [f"{pad}<h3>{esc(lines(block[1])[block[2]][1])}</h3>"]
    if kind == "tree":                   # ("tree", slide, start, end) -> source paragraph levels
        return r_list(lines(block[1])[block[2]:block[3]], pad)
    if kind == "treed":                  # ("treed", slide, [(index, depth), ...]) -> explicit depths
        return r_list([(d, lines(block[1])[i][1]) for i, d in block[2]], pad)
    if kind == "code":                   # ("code", slide, start, end, title[, prefix lines])
        prefix = list(block[5]) if len(block) > 5 else []
        return r_code(prefix + [t for _, t in lines(block[1])[block[2]:block[3]]], pad, block[4])
    if kind == "img":                    # ("img", [(slide, idx, width, alt), ...])
        return r_images(block[1], pad)
    if kind == "video":                  # ("video", id, summary, watch_url[, start])
        return r_video(block[1], block[2], block[3], pad, *(block[4:5] or [None]))
    if kind == "table":                  # ("table", slide)
        return r_table(block[1], pad)
    if kind == "link":                   # ("link", label, url)
        return [f"{pad}<p>{link(block[2], block[1])}</p>"]
    if kind == "refs":                   # ("refs", summary, [(label, url), ...])
        items = [f"{pad}    <ul>"] + [f"{pad}      <li>{link(u, l)}</li>" for l, u in block[2]] + [f"{pad}    </ul>"]
        return r_nested(block[1], items, pad)
    if kind == "tut":                    # ("tut", page, label)
        return [f'{pad}<p><a href="../index.html?page=pages/{block[1]}" target="_top">{esc(block[2])}</a></p>']
    if kind == "from":                   # ("from", slide, index) -> "From <url>" with url isolated for linkify
        m = re.fullmatch(r"(.*?<)(https?://\S+?)(>)", lines(block[1])[block[2]][1])
        return [f"{pad}<p>{esc(m.group(1))}<span>{esc(m.group(2))}</span>{esc(m.group(3))}</p>"]
    if kind == "notes":                  # ("notes", slide[, start, end]) -> speaker notes accordion
        paras = [re.sub(r"\s+", " ", t).strip() for t in SLIDES[block[1]].notes_slide.notes_text_frame.text.split("\n")]
        paras = [t for t in paras if t][(block[2] if len(block) > 2 else None):(block[3] if len(block) > 3 else None)]
        return r_nested("Speaker notes", [f"{pad}    <p>{esc(t)}</p>" for t in paras], pad)
    if kind == "notecode":               # ("notecode", slide, title) -> speaker-notes code as a code panel
        code = [re.sub(r"\s+", " ", t).strip() for t in SLIDES[block[1]].notes_slide.notes_text_frame.text.split("\n")]
        return r_code([t for t in code if t], pad, block[2])
    if kind == "codelist":               # ("codelist", [code lines], title)
        return r_code(block[1], pad, block[2])
    if kind == "nested":                 # ("nested", summary, [blocks])
        inner = []
        for b in block[2]:
            inner.extend(render_block(b, pad + "    "))
        return r_nested(block[1], inner, pad)
    raise ValueError(kind)


def heading(num, idx=0):
    return lines(num)[idx][1]




YT = "https://www.youtube.com/watch?v="
NOORCON = "https://universityofwestminster-my.sharepoint.com/:b:/g/personal/fergusj_westminster_ac_uk/EYneh4DsN09JrpgJO040U64B0j2ChZPEx1XezbrDF5r1hA?e=RhlZ4c"
STACK = "https://stackoverflow.com/questions/61417584/unity-player-controller-not-moving-toward-mouse-cursor"


def img(num, idx, width, alt):
    return (num, idx, width, alt)


def merged(num, start, end, continuations):
    """Code lines with wrapped continuation lines re-joined onto the line above."""
    out = []
    for i in range(start, end):
        text = lines(num)[i][1]
        if i in continuations:
            out[-1] = out[-1] + " " + text
        else:
            out.append(text)
    return out


def sub(num, depth_by_index):
    return ("treed", num, depth_by_index)


# ---------- consolidation plan: 55 source slides -> 14 themed slides ----------
# section: dict(s=summary, src=[source slides], b=[blocks])
PLAN = [
    {"h2": "5MMCS001W Lecture 11 - Cameras and Lerp", "src": "1-4", "sections": [
        {"s": "Lecture Overview", "src": [1, 2], "b": [
            ("p", 1, 0, 1),
            ("p", 2, 0, 1),
        ]},
        {"s": heading(3) + " - " + heading(3, 1), "src": [3], "b": [
            ("p", 3, 2, 4),
        ]},
        {"s": "Camera Prefabs Folder", "src": [4], "b": [
            ("p", 4, 1, 5),
            ("img", [img(4, 0, "45%", "Unity Project window: Assets > Prefabs with Camera Align with View, Camera RTS and Main Camera prefabs"),
                     img(4, 1, "45%", "Unity Project window: Standard Assets > Cameras > Prefabs with CctvCamera, FreeLookCameraRig, HandheldCamera and MultipurposeCameraRig")]),
        ]},
        {"s": "Related Tutorials", "src": [], "b": [
            ("tut", "tut-02-03-students.html", "Tutorial 02-03: Unity Intro - Interaction Scripting"),
            ("tut", "tut-04-students.html", "Tutorial 04: Unity Animation and Animation Import Overview"),
            ("tut", "tut-08-character-troubleshooting-bb.html", "Tutorial 08: Unity - 3 types of character controllers"),
        ]},
    ]},
    {"h2": "Camera and Character Controller Prefabs", "src": "5-9", "sections": [
        {"s": heading(5), "src": [5], "b": [
            ("video", "z_mJ_BjFqlU", "Unity 3D Tutorial Third Person Shooter part 2 The Aiming [Outdated]", YT + "z_mJ_BjFqlU"),
            ("p", 5, 2, 8),
        ]},
        {"s": heading(6), "src": [6], "b": [
            ("p", 6, 1, 6),
        ]},
        {"s": heading(7), "src": [7], "b": [
            ("p", 7, 1, 5),
        ]},
        {"s": heading(8), "src": [8], "b": [
            ("p", 8, 1, 2),
            ("tree", 8, 2, 4),
        ]},
        {"s": heading(9), "src": [9], "b": [
            ("p", 9, 1, 4),
            ("tree", 9, 4, 6),
            ("img", [img(9, 0, "48%", "Unity Asset Store page: Starter Assets - ThirdPerson, Updates in new CharacterController package"),
                     img(9, 1, "48%", "Unity Asset Store page: Starter Assets - FirstPerson, Updates in new CharacterController package")]),
        ]},
    ]},
    {"h2": heading(10), "src": "10-12", "sections": [
        {"s": heading(10, 2), "src": [10], "b": [
            ("video", "dXzF-S9tVrM", "Third Person Character Controller Setup in Unity", YT + "dXzF-S9tVrM"),
            sub(10, [(3, 0), (4, 0), (5, 1), (6, 1), (7, 0), (8, 0), (9, 1), (10, 0), (11, 0), (12, 1), (13, 1), (14, 0), (15, 1), (16, 1)]),
        ]},
        {"s": heading(11, 2), "src": [11], "b": [
            sub(11, [(3, 0), (4, 0), (5, 0), (6, 0), (7, 1), (8, 1), (9, 0)]),
        ]},
        {"s": heading(12), "src": [12], "b": [
            ("p", 12, 2, 3),
            ("tree", 12, 3, 15),
        ]},
    ]},
    {"h2": "Camera Controllers - Other Scripts and SmoothFollow", "src": "13-14", "sections": [
        {"s": heading(13), "src": [13], "b": [
            ("p", 13, 1, 2),
            sub(13, [(2, 0), (3, 0), (4, 0), (5, 0), (6, 0), (7, 0)]),
        ]},
        {"s": heading(14), "src": [14], "b": [
            ("p", 14, 1, 4),
            ("code", 14, 4, 17, "EcameraSmoothFollow.cs"),
        ]},
    ]},
    {"h2": "Unity Camera Prefabs - FreeLook", "src": "15-16", "sections": [
        {"s": heading(15), "src": [15], "b": [
            ("p", 15, 1, 3),
            sub(15, [(3, 0), (4, 1), (5, 0), (6, 0), (7, 0), (8, 0)]),
            ("img", [img(15, 0, "45%", "Unity Project window: Standard Assets > Cameras > Prefabs showing FreeLookCameraRig")]),
        ]},
        {"s": heading(16), "src": [16], "b": [
            ("p", 16, 1, 5),
            sub(16, [(5, 0), (6, 1), (7, 2)]),
        ]},
    ]},
    {"h2": heading(17), "src": "17-19", "sections": [
        {"s": heading(18, 1), "src": [17, 18], "b": [
            ("p", 17, 1, 3),
            ("tree", 18, 2, 5),
        ]},
        {"s": heading(19, 2), "src": [19], "b": [
            ("p", 19, 0, 2),
            ("video", "8HALMRhbpGI", "3D Menu + Camera Waypoints", YT + "8HALMRhbpGI"),
            ("p", 19, 4, 11),
            ("refs", "Camera documentation links", [
                ("Camera.WorldToViewportPoint", "http://docs.unity3d.com/ScriptReference/Camera.WorldToViewportPoint.html"),
                ("Camera.ViewportToScreenPoint", "http://docs.unity3d.com/ScriptReference/Camera.ViewportToScreenPoint.html"),
            ]),
        ]},
    ]},
    {"h2": heading(20), "src": "20-22", "sections": [
        {"s": heading(20, 2) + " - On your main camera", "src": [20], "b": [
            ("p", 20, 1, 3),
            ("from", 20, 3),
            ("p", 20, 4, 5),
            ("code", 20, 5, 13, "FocusScript.cs - Focus()"),
        ]},
        {"s": heading(21, 2) + " - On your GameObject", "src": [21], "b": [
            ("p", 21, 4, 5),
            ("code", 21, 5, 13, "OnMouseOver() on the GameObject"),
        ]},
        {"s": heading(22, 2), "src": [22], "b": [
            ("p", 22, 1, 2),
            ("p", 22, 3, 11),
        ]},
    ]},
    {"h2": "Unity Camera Controllers - RTS Camera Movement", "src": "23-30", "sections": [
        {"s": heading(23), "src": [23], "b": [
            ("p", 23, 1, 5),
        ]},
        {"s": heading(24, 2), "src": [24, 25], "b": [
            ("p", 24, 0, 2),
            ("p", 24, 3, 7),
            ("tree", 24, 7, 13),
            ("codelist", merged(25, 4, 15, {7, 11}), "RTS Position - Basic with RMB"),
        ]},
        {"s": heading(26, 2), "src": [26, 27], "b": [
            ("p", 26, 0, 2),
            ("p", 26, 3, 7),
            ("code", 27, 0, 20, "RTS Position - Mouse 2D Out Screen Drag (script on Camera)"),
        ]},
        {"s": heading(28, 2), "src": [28, 29], "b": [
            ("p", 28, 0, 2),
            ("p", 28, 3, 7),
            ("code", 29, 0, 20, "RTS Position - Mouse 2D In Screen Drag (script on Camera)"),
        ]},
        {"s": heading(30), "src": [30], "b": [
            ("p", 30, 1, 5),
        ]},
    ]},
    {"h2": heading(31), "src": "31-36", "sections": [
        {"s": heading(32), "src": [31, 32], "b": [
            ("p", 31, 1, 2),
            ("p", 32, 1, 3),
            ("video", "CvEMvYUXxaE", "Testing a dynamic 2.5D camera on Unity 3D engine", YT + "CvEMvYUXxaE"),
            ("p", 32, 3, 4),
            ("link", heading(32, 4), NOORCON),
            sub(32, [(5, 0), (6, 1), (7, 0), (8, 0)]),
        ]},
        {"s": heading(33), "src": [33], "b": [
            ("p", 33, 3, 4),
            ("tree", 33, 4, 6),
        ]},
        {"s": "Diorama Tutorial - playerInput Script", "src": [34], "b": [
            ("p", 34, 2, 3),
            ("code", 34, 3, 28, "playerInput.cs"),
        ]},
        {"s": "Diorama Tutorial - Scene Setup and Align With View", "src": [35], "b": [
            ("p", 35, 3, 4),
            ("tree", 35, 4, 8),
            ("p", 35, 8, 11),
        ]},
        {"s": "Diorama Tutorial - Static Colliders and Camera Alignment", "src": [36], "b": [
            ("p", 36, 3, 4),
            ("tree", 36, 4, 5),
            ("p", 36, 5, 8),
        ]},
    ]},
    {"h2": "Camera Switcher and Camera Exercises", "src": "37-41", "sections": [
        {"s": heading(37), "src": [37], "b": [
            ("p", 37, 1, 4),
        ]},
        {"s": "Exercise: Unity 3D Camera - Waypoints and Mouse Orbits", "src": [38], "b": [
            ("p", 38, 0, 2),
            ("tree", 38, 2, 7),
            ("refs", "Speaker notes links", [
                ("YouTube video -5RiMoILHTA (currently unavailable)", "https://www.youtube.com/watch?v=-5RiMoILHTA"),
                ("Unity forum: Fly/move camera object to waypoint with button click", "https://forum.unity3d.com/threads/fly-move-camera-object-to-waypoint-with-button-click.46988/"),
            ]),
        ]},
        {"s": heading(39), "src": [39], "b": [
            ("p", 39, 1, 2),
            sub(39, [(2, 0), (3, 0), (4, 1), (5, 0), (6, 1), (7, 0), (8, 1)]),
        ]},
        {"s": heading(40), "src": [40], "b": [
            sub(40, [(1, 0), (2, 1), (3, 0), (4, 1)]),
        ]},
        {"s": heading(41), "src": [41], "b": [
            sub(41, [(1, 0), (2, 1), (3, 1), (4, 0), (5, 1)]),
        ]},
    ]},
    {"h2": heading(42), "src": "42-47", "sections": [
        {"s": "LERP and Coroutine References", "src": [42], "b": [
            ("p", 42, 1, 2),
            ("video", "RNccTrsgO9g", "The right way to Lerp in Unity", "https://youtu.be/RNccTrsgO9g"),
            ("p", 42, 3, 4),
            ("video", "kUP6OK36nrM", "Coroutines in Unity (how & when to use them)", "https://youtu.be/kUP6OK36nrM"),
        ]},
        {"s": heading(43, 1), "src": [43], "b": [
            ("tree", 43, 2, 6),
            ("img", [img(43, 0, "48%", "Lerp in Update: lerpedValue from start to end over duration using timeElapsed / duration"),
                     img(43, 1, "48%", "Lerp in a coroutine: IEnumerator LerpValue(start, end) started from Start()")]),
        ]},
        {"s": heading(44, 1).strip(" -"), "src": [44], "b": [
            ("pj", 44, 1, 3),
            ("tree", 44, 3, 9),
            ("img", [img(44, 0, "55%", "Dynamic easing: lerpedValue = Mathf.Lerp(lerpedValue, end, Time.deltaTime) in Update"),
                     img(44, 1, "40%", "Square easing towards the mouse cursor: position = Vector2.Lerp(position, target, time)")]),
        ]},
        {"s": heading(45, 1), "src": [45], "b": [
            ("pj", 45, 2, 4),
            sub(45, [(4, 0), (5, 0), (6, 0), (7, 1), (8, 1), (9, 1), (10, 1)]),
            ("img", [img(45, 0, "50%", "Mathf.SmoothDamp(target, velocity, smooth time)"),
                     img(45, 1, "60%", "SmoothDamp in Update: lerpedValue = Mathf.SmoothDamp(lerpedValue, 1, ref velocity, duration)")]),
        ]},
        {"s": heading(46, 1), "src": [46], "b": [
            ("p", 46, 2, 3),
            ("img", [img(46, 0, "50%", "Lerp coroutine with ease in and out: t = t * t * (3f - 2f * t)")]),
        ]},
        {"s": heading(47, 1), "src": [47], "b": [
            ("img", [img(47, 0, "45%", "Unity animation curve editor with an ease curve"),
                     img(47, 1, "50%", "Lerp coroutine using t = animCurve.Evaluate(t) from a public AnimationCurve")]),
        ]},
    ]},
    {"h2": "ClickLerp: New Camera Position by Keypress and Mouse", "src": "48-49", "sections": [
        {"s": heading(48), "src": [48], "b": [
            ("code", 48, 1, 4, "OnMouseEnter()"),
            ("p", 48, 4, 5),
            ("tree", 48, 5, 8),
            ("p", 48, 8, 9),
        ]},
        {"s": heading(49), "src": [49], "b": [
            ("code", 49, 1, 19, "ClickLerp.cs"),
        ]},
        {"s": "Speaker notes link (slides 48-54)", "src": [], "b": [
            ("link", "Stack Overflow: Unity player controller not moving toward mouse cursor", STACK),
        ]},
    ]},
    {"h2": heading(50), "src": "50-54", "sections": [
        {"s": "Player Game Object - PlayerMovement.cs", "src": [50], "b": [
            sub(50, [(1, 0), (2, 1)]),
            ("p", 50, 3, 4),
            sub(50, [(4, 0), (5, 0)]),
            ("p", 50, 6, 7),
            ("code", 50, 7, 9, "PlayerMovement.cs - input"),
            ("pj", 50, 9, 11),
            ("code", 50, 11, 15, "PlayerMovement.cs - FixedUpdate movement"),
            ("p", 50, 15, 16),
            ("img", [img(50, 0, "40%", "Unity Inspector for the Player GameObject with Rigidbody, Capsule Collider and PlayerMovement script")]),
        ]},
        {"s": "Player Game Object - Orientation Game Object", "src": [51], "b": [
            sub(51, [(1, 0), (2, 1)]),
            ("p", 51, 3, 4),
            ("img", [img(51, 0, "45%", "Unity Hierarchy and Inspector with the Orientation GameObject selected under Player")]),
        ]},
        {"s": "Player Game Object - Camera Position Game Object", "src": [52], "b": [
            sub(52, [(1, 0), (2, 1)]),
            ("p", 52, 3, 4),
            ("img", [img(52, 0, "45%", "Unity Hierarchy and Inspector with the CameraPosition GameObject selected under Player")]),
        ]},
        {"s": "Camera Holder Game Object - CameraMove.cs", "src": [53], "b": [
            sub(53, [(1, 0), (2, 1)]),
            ("pj", 53, 3, 5),
            ("pj", 53, 5, 7),
            ("pj", 53, 7, 9),
            ("p", 53, 9, 10),
            ("img", [img(53, 0, "40%", "Unity Inspector for the CameraHolder GameObject with the CameraMove script")]),
        ]},
        {"s": "Camera Holder Game Object - PlayerCamera - PlayerCam.cs", "src": [54], "b": [
            sub(54, [(1, 0), (2, 1), (3, 2)]),
            ("pj", 54, 4, 7),
            ("pj", 54, 7, 11),
            ("pj", 54, 11, 13),
            ("p", 54, 13, 14),
            ("img", [img(54, 0, "40%", "Unity Inspector for the PlayerCamera GameObject with Camera, Audio Listener and PlayerCam script")]),
        ]},
    ]},
    {"h2": "Unity 2.5D Camera - Position State Changes", "src": "55", "sections": [
        {"s": heading(55), "src": [55], "b": [
            ("p", 55, 1, 3),
            sub(55, [(3, 0), (4, 0)]),
            ("video", "z_mJ_BjFqlU", heading(55, 5), YT + "z_mJ_BjFqlU", 720),
            sub(55, [(6, 0), (7, 1)]),
            ("p", 55, 8, 9),
            ("notecode", 55, "Speaker notes code: DOMove to clicked position"),
        ]},
    ]},
]


def render_section(sec, number, sub, first):
    pad = "          "
    body = []
    for block in sec["b"]:
        body.extend(render_block(block, pad))
    src = ", ".join(str(n) for n in sec["src"]) or "added cross-reference"
    return [
        f"      <!-- Section {number}.{sub}: {sec['s']} (source slide {src}) -->",
        f'      <details class="accordion"{" open" if first else ""}>',
        f"        <summary>{esc(sec['s'])}</summary>",
        '        <div class="accordion-body">',
        *body,
        "        </div>",
        "      </details>",
    ]


def render_slide(spec, number, total):
    first = number == 1
    out = [
        f"    <!-- Slide {number}: {spec['h2']} (source slides {spec['src']}) -->",
        f'    <div class="slide{" active" if first else ""}">',
        f'      <div class="slide-title">{SLIDE_TITLE}</div>',
        '      <div class="slide-content slide-content-accordion">',
        f"        <h2>{esc(spec['h2'])}</h2>",
        '        <div class="accordion-scroll">',
    ]
    for i, sec in enumerate(spec["sections"]):
        out.extend(render_section(sec, number, i + 1, i == 0))
    out += [
        "        </div>",
        "      </div>",
        '      <div class="slide-controls">',
        f'        <button class="slide-button" data-slide-step="-1"{" disabled" if first else ""}> ← Previous</button>',
        '        <img src="images/UOW_Logo_Length_Alpha.png" class="slide-logo" alt="University of Westminster" />',
        f'        <button class="slide-button" data-slide-step="1"{" disabled" if number == total else ""}>Next → </button>',
        f'        <div class="slide-counter"><span id="current-slide-{number}">{number}</span>'
        f' / <span id="total-slides">{total}</span></div>',
        "      </div>",
        '      <div class="keyboard-hint">',
        "        Use <kbd>→</kbd> <kbd>←</kbd> or click buttons to navigate",
        "      </div>",
        "    </div>",
        "",
    ]
    return out


def main():
    covered = {n for spec in PLAN for sec in spec["sections"] for n in sec["src"]}
    missing = set(SLIDES) - covered
    if missing:
        raise SystemExit(f"Source slides not mapped: {sorted(missing)}")

    original = FOOTER_SOURCE.read_text(encoding="utf-8").split("\n")
    marker = '  <script src="../assets/js/site.js'
    footer = original[next(i for i, ln in enumerate(original) if ln.startswith(marker)):]

    total = len(PLAN)
    doc = [
        "<!DOCTYPE html>",
        '<html lang="en">',
        "<head>",
        '  <meta charset="UTF-8" />',
        '  <meta name="viewport" content="width=device-width, initial-scale=1.0" />',
        f"  <title>{PAGE_TITLE}</title>",
        '  <link rel="preconnect" href="https://fonts.googleapis.com" />',
        '  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />',
        '  <link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;600;700&family=Space+Grotesk:wght@500;700&display=swap" rel="stylesheet" />',
        '  <link rel="stylesheet" href="../assets/css/site.css?v=20261007a" />',
        '  <link rel="stylesheet" href="../assets/css/slideshow.css?v=20261007a" />',
        "</head>",
        '<body class="lecture-fullbleed">',
        '  <div class="slideshow-container">',
    ]
    for i, spec in enumerate(PLAN, 1):
        doc.extend(render_slide(spec, i, total))
    doc += ["  </div>", ""] + footer

    OUT_PATH.write_text("\n".join(doc), encoding="utf-8")
    print(f"wrote {OUT_PATH} with {total} slides from {len(SLIDES)} source slides")


if __name__ == "__main__":
    main()
